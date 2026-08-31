import sys
import types
import unittest
from datetime import date
from decimal import Decimal
from pathlib import Path
from types import SimpleNamespace
from unittest.mock import AsyncMock, patch


sys.path.insert(0, str(Path(__file__).resolve().parents[1] / "api"))

# CI для чистих unit-тестів може не мати важких runtime-залежностей API.
try:
    import fastapi  # noqa: F401
except ModuleNotFoundError:
    fastapi = types.ModuleType("fastapi")

    class _APIRouter:
        def __init__(self, *args, **kwargs):
            pass

        def _route(self, *args, **kwargs):
            return lambda fn: fn

        get = post = put = patch = delete = _route

    class _HTTPException(Exception):
        def __init__(self, status_code, detail):
            super().__init__(detail)
            self.status_code = status_code
            self.detail = detail

    fastapi.APIRouter = _APIRouter
    fastapi.HTTPException = _HTTPException
    fastapi.Query = lambda default=None, **kwargs: default
    responses = types.ModuleType("fastapi.responses")
    responses.StreamingResponse = object
    sys.modules["fastapi"] = fastapi
    sys.modules["fastapi.responses"] = responses

try:
    import pydantic  # noqa: F401
except ModuleNotFoundError:
    pydantic = types.ModuleType("pydantic")
    pydantic.BaseModel = object
    pydantic.Field = lambda default_factory=None, **kwargs: default_factory()
    sys.modules["pydantic"] = pydantic

try:
    import openpyxl  # noqa: F401
except ModuleNotFoundError:
    openpyxl = types.ModuleType("openpyxl")
    openpyxl.Workbook = object
    styles = types.ModuleType("openpyxl.styles")
    styles.Alignment = styles.Font = styles.PatternFill = object
    sys.modules["openpyxl"] = openpyxl
    sys.modules["openpyxl.styles"] = styles

from app import fuel  # noqa: E402
from fastapi import HTTPException  # noqa: E402


class _Transaction:
    async def __aenter__(self):
        return self

    async def __aexit__(self, exc_type, exc, tb):
        return False


class _Acquire:
    def __init__(self, conn):
        self.conn = conn

    async def __aenter__(self):
        return self.conn

    async def __aexit__(self, exc_type, exc, tb):
        return False


class _FakeConn:
    def __init__(self, existing=None, blocking=None):
        self.existing = existing
        self.blocking = blocking
        self.executed = []

    def transaction(self):
        return _Transaction()

    async def fetchrow(self, query, *args):
        if "FROM vehicles" in query:
            return {"id": 9, "name": "PEUGEOT", "plate": "КА6843PI"}
        if "FROM drivers" in query:
            return {"id": 27, "name": "Новий водій"}
        if "SELECT id,status FROM transport_sheets" in query:
            return self.existing
        if "work_date>$2 AND status='approved'" in query:
            return self.blocking
        if "INSERT INTO transport_sheets" in query:
            return {"id": 501, "vehicle_id": 9, "driver_id": 27,
                    "work_date": date(2026, 8, 31), "status": "submitted"}
        raise AssertionError(f"Unexpected query: {query}")

    async def execute(self, query, *args):
        self.executed.append((query, args))


class _FakePool:
    def __init__(self, conn):
        self.conn = conn

    def acquire(self):
        return _Acquire(self.conn)


def _body(**changes):
    values = {
        "work_date": date(2026, 8, 31),
        "vehicle_id": 9,
        "driver_id": 27,
        "odometer_start": "125420,0",
        "odometer_end": "125518,0",
        "refuels": [SimpleNamespace(liters="35,5"), SimpleNamespace(liters=10)],
        "reason": "Водій працював без робочого телефону",
    }
    values.update(changes)
    return SimpleNamespace(**values)


class LogistCreateSheetTest(unittest.IsolatedAsyncioTestCase):
    async def test_creates_submitted_sheet_with_driver_odometer_and_refuels(self):
        conn = _FakeConn()
        old_pool = fuel.pool
        fuel.pool = _FakePool(conn)
        self.addCleanup(setattr, fuel, "pool", old_pool)
        with patch.object(fuel, "_derive_opening", AsyncMock(return_value=Decimal("18.25"))), \
             patch.object(fuel, "_recalc_chain", AsyncMock()) as recalc:
            result = await fuel.create_sheet_by_logist(_body())

        self.assertEqual(result, {"ok": True, "sheet_id": 501})
        inserts = [args for query, args in conn.executed
                   if "INSERT INTO transport_sheet_refuels" in query]
        self.assertEqual([x[1] for x in inserts], [Decimal("35.5"), Decimal("10")])
        recalc.assert_awaited_once_with(
            9, date(2026, 8, 31), "Створення листа логістом")

    async def test_rejects_duplicate_vehicle_date(self):
        conn = _FakeConn(existing={"id": 77, "status": "submitted"})
        old_pool = fuel.pool
        fuel.pool = _FakePool(conn)
        self.addCleanup(setattr, fuel, "pool", old_pool)
        with patch.object(fuel, "_derive_opening", AsyncMock()), \
             patch.object(fuel, "_recalc_chain", AsyncMock()):
            with self.assertRaises(HTTPException) as caught:
                await fuel.create_sheet_by_logist(_body(refuels=[]))
        self.assertEqual(caught.exception.status_code, 409)
        self.assertIn("№77", caught.exception.detail)

    async def test_rejects_insert_before_later_approved_sheet(self):
        conn = _FakeConn(blocking={"work_date": date(2026, 9, 2)})
        old_pool = fuel.pool
        fuel.pool = _FakePool(conn)
        self.addCleanup(setattr, fuel, "pool", old_pool)
        with patch.object(fuel, "_derive_opening", AsyncMock()), \
             patch.object(fuel, "_recalc_chain", AsyncMock()):
            with self.assertRaises(HTTPException) as caught:
                await fuel.create_sheet_by_logist(_body(refuels=[]))
        self.assertEqual(caught.exception.status_code, 409)
        self.assertIn("2026-09-02", caught.exception.detail)

    async def test_rejects_decreasing_odometer_before_db_write(self):
        with self.assertRaises(HTTPException) as caught:
            await fuel.create_sheet_by_logist(
                _body(odometer_start=200, odometer_end=199, refuels=[]))
        self.assertEqual(caught.exception.status_code, 400)


if __name__ == "__main__":
    unittest.main()
